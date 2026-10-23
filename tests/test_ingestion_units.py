import io

import pytest

from lumen.adapters.embedding import MockEmbeddingProvider
from lumen.adapters.storage import LocalDiskStorage
from lumen.ingestion.chunker import StructuralChunker
from lumen.ingestion.extract import (
    DOCX,
    DocxExtractor,
    ExtractorRegistry,
    PdfExtractor,
    TextExtractor,
    serialize_tables_to_markdown,
    synthetic_paginate,
)
from lumen.ingestion.headers import ContextualHeaderGenerator
from lumen.ingestion.pipeline import RETRY_POLICIES
from lumen.types import ChunkCandidate
from lumen.utils import TokenBucketLimiter, cosine_similarity

chunker = StructuralChunker()


class TestStructuralChunker:
    def test_short_page_is_one_chunk_with_heading_as_section(self):
        chunks = chunker.chunk_page("# Section 4: Limitation of Liability\n\nSection 4.2 does not apply in this case.")
        assert len(chunks) == 1
        section, text, _ = chunks[0]
        assert section == "Section 4: Limitation of Liability"
        assert "Section 4.2 does not apply" in text

    def test_never_splits_a_markdown_table(self):
        table = "\n".join(["| A | B |", "| --- | --- |", "| 1 | 2 |", "| 3 | 4 |"])
        chunks = chunker.chunk_page(f"# Pricing\n\nSome intro text.\n\n{table}\n\nSome trailing text.")
        holder = next(text for _, text, _ in chunks if "| 1 | 2 |" in text)
        assert "| 3 | 4 |" in holder

    def test_oversized_paragraph_is_split_without_empty_chunks(self):
        big = " ".join(f"This is sentence number {i} in a very long paragraph." for i in range(400))
        chunks = chunker.chunk_page(big)
        assert len(chunks) > 1
        assert all(text.strip() for _, text, _ in chunks)
        assert all(tokens <= 800 + 64 for _, _, tokens in chunks)

    def test_consecutive_chunks_overlap(self):
        big = " ".join(f"Sentence {i} with some additional filler words to reach target length." for i in range(300))
        chunks = chunker.chunk_page(big)
        assert len(chunks) >= 2
        tail = chunks[0][1].split(" ")[-3:]
        assert " ".join(tail) in chunks[1][1]

    def test_list_is_kept_together(self):
        page = "# Steps\n\n- first step\n- second step\n- third step\n\nAfter."
        holder = next(text for _, text, _ in chunker.chunk_page(page) if "- first step" in text)
        assert "- third step" in holder


class TestExtraction:
    def test_table_heuristic_turns_columnar_lines_into_markdown(self):
        text = "Plan    Price    Seats\nStarter    $99    5\nGrowth    $349    25\nnot a row"
        md = serialize_tables_to_markdown(text)
        assert "| Plan | Price | Seats |" in md and "| --- | --- | --- |" in md and "| Growth | $349 | 25 |" in md
        assert md.endswith("not a row")

    def test_two_columnar_lines_are_not_a_table(self):
        assert serialize_tables_to_markdown("a    b\nc    d") == "a    b\nc    d"

    def test_synthetic_pagination_respects_paragraph_boundaries(self):
        paras = ["x" * 1000 for _ in range(7)]
        pages = synthetic_paginate("\n\n".join(paras))
        # Two 1000-char paragraphs fit a 3000-char page; a third would overflow it.
        assert [p.page_number for p in pages] == [1, 2, 3, 4]
        assert all(len(p.text) <= 3000 for p in pages)
        assert all(not p.text.startswith("\n") for p in pages)

    def test_markdown_outline_lists_headings(self):
        result = TextExtractor().extract(b"# Title\n\nbody\n\n## Section 2\n\nmore")
        assert result.outline == "# Title\n## Section 2"

    def test_docx_headings_lists_and_tables_become_markdown(self):
        from docx import Document

        doc = Document()
        doc.add_heading("Pricing Policy", level=1)
        doc.add_paragraph("Intro paragraph.")
        doc.add_paragraph("first item", style="List Bullet")
        doc.add_paragraph("second item", style="List Bullet")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Plan", "Price"
        table.cell(1, 0).text, table.cell(1, 1).text = "Growth", "$349"
        buf = io.BytesIO()
        doc.save(buf)

        result = DocxExtractor().extract(buf.getvalue())
        text = result.pages[0].text
        assert text.startswith("# Pricing Policy")
        assert "- first item\n- second item" in text
        assert "| Plan | Price |\n| --- | --- |\n| Growth | $349 |" in text
        assert result.outline == "# Pricing Policy"

    def test_pdf_pages_are_extracted_separately(self, tiny_pdf):
        result = PdfExtractor().extract(tiny_pdf)
        assert [p.page_number for p in result.pages] == [1, 2]
        assert "Limitation of Liability" in result.pages[0].text
        assert "Governing law" in result.pages[1].text

    def test_registry_rejects_unsupported_types(self):
        with pytest.raises(ValueError, match="PPTX"):
            ExtractorRegistry().extract(
                b"", "application/vnd.openxmlformats-officedocument.presentationml.presentation"
            )
        with pytest.raises(ValueError, match="Unsupported"):
            ExtractorRegistry().extract(b"", "image/png")
        assert DOCX in [m for m in ("application/pdf", DOCX)]


async def test_mock_embeddings_are_deterministic_unit_vectors():
    provider = MockEmbeddingProvider(64)
    a1, a2, b = await provider.embed(["alpha", "alpha", "beta"])
    assert a1 == a2
    assert abs(cosine_similarity(a1, a1) - 1) < 1e-9
    assert abs(cosine_similarity(a1, b)) < 0.5


async def test_header_generator_falls_back_to_deterministic_headers():
    class BrokenCompletion:
        model = "broken"

        async def complete(self, *args, **kwargs):
            return "this is not json", None

    candidates = [ChunkCandidate(3, "Section 4", "text", 1), ChunkCandidate(5, "", "text", 1)]
    headers = await ContextualHeaderGenerator(BrokenCompletion(), use_llm=True).generate("Acme MSA", "", candidates)
    assert headers == ["Acme MSA - Section 4", "Acme MSA - page 5"]


async def test_local_storage_refuses_path_traversal(tmp_path):
    storage = LocalDiskStorage(str(tmp_path))
    await storage.put("ws/doc.txt", b"hi", "text/plain")
    assert await storage.get("ws/doc.txt") == b"hi"
    with pytest.raises(ValueError):
        await storage.get("../outside.txt")


async def test_token_bucket_throttles_beyond_capacity():
    import time

    limiter = TokenBucketLimiter(capacity=5, refill_per_second=50)
    start = time.monotonic()
    await limiter.acquire(5)
    await limiter.acquire(5)
    assert time.monotonic() - start >= 0.08


def test_retry_policies_match_stage_failure_modes():
    assert RETRY_POLICIES["extract"].attempts == 2 and RETRY_POLICIES["extract"].backoff == "fixed"
    assert RETRY_POLICIES["embed"].attempts == 5
    assert RETRY_POLICIES["embed"].delay_for(1) == 4 and RETRY_POLICIES["embed"].delay_for(3) == 16
    assert RETRY_POLICIES["extract"].delay_for(2) == 5


def _tesseract_available() -> bool:
    import os
    import shutil

    return bool(shutil.which("tesseract") or os.environ.get("TESSERACT_CMD"))


@pytest.mark.skipif(not _tesseract_available(), reason="tesseract binary not installed")
def test_scanned_pdf_page_falls_back_to_ocr():
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("RGB", (1400, 400), "white")
    try:
        font = ImageFont.truetype("arial.ttf", 64)
    except OSError:
        font = ImageFont.load_default(size=64)
    ImageDraw.Draw(image).text((40, 150), "Invoice total due 12450 dollars", fill="black", font=font)
    buf = io.BytesIO()
    image.save(buf, format="PDF")

    result = PdfExtractor().extract(buf.getvalue())
    assert result.pages[0].ocr is True
    assert "12450" in result.pages[0].text
