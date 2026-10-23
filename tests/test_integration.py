"""End-to-end tests against real MongoDB and Redis: the FastAPI app is driven over ASGI and the
ingestion pipeline runs inline with the same stage code and retry policies the arq workers use."""

import json

import httpx
import pytest
import pytest_asyncio

from lumen.api.app import create_app
from lumen.composition import build_container
from lumen.ingestion.extract import DOCX
from lumen.query.core import REFUSAL_MESSAGE

pytestmark = pytest.mark.integration

MSA = (
    b"# Acme Master Services Agreement 2024\n\n## Section 4: Limitation of Liability\n\n"
    b"Section 4.2 does not apply in the case of gross negligence or willful misconduct.\n\n"
    b"## Section 7: Termination\n\nEarly cancellation requires a pro-rated refund of prepaid fees."
)


@pytest_asyncio.fixture(scope="session")
async def container(integration_settings):
    c = await build_container(integration_settings, queue="inline")
    await c.redis.flushdb()
    yield c
    await c.mongo.drop_database(c.settings.mongo_db_name)
    await c.redis.flushdb()
    await c.close()


@pytest_asyncio.fixture(scope="session")
async def client(container):
    app = create_app(container=container)
    app.state.container = container
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as c:
        yield c


async def signup(client, email: str, workspace: str) -> dict:
    res = await client.post(
        "/auth/signup", json={"email": email, "password": "password123", "workspaceName": workspace}
    )
    assert res.status_code == 201, res.text
    return res.json()


def auth(session: dict) -> dict:
    return {"Authorization": f"Bearer {session['token']}"}


async def ask(client, session: dict, question: str, **filters) -> dict[str, list]:
    events: dict[str, list] = {}
    body = {"question": question, **({"filters": filters} if filters else {})}
    async with client.stream("POST", "/query", json=body, headers=auth(session)) as res:
        assert res.status_code == 200
        assert res.headers["content-type"].startswith("text/event-stream")
        raw = (await res.aread()).decode()
    for frame in raw.strip().split("\n\n"):
        event_line, data_line = frame.split("\n", 1)
        events.setdefault(event_line.removeprefix("event: "), []).append(json.loads(data_line.removeprefix("data: ")))
    return events


async def upload(client, session: dict, name: str, data: bytes, mime: str) -> dict:
    res = await client.post("/documents", files={"file": (name, data, mime)}, headers=auth(session))
    assert res.status_code == 201, res.text
    return res.json()


async def test_auth_flow_and_jwt_scoping(client):
    s = await signup(client, "owner@acme.test", "Acme Legal")
    assert s["user"]["role"] == "admin" and s["workspace"]["slug"] == "acme-legal"

    assert (
        await client.post(
            "/auth/signup", json={"email": "owner@acme.test", "password": "password123", "workspaceName": "Dup"}
        )
    ).status_code == 409
    assert (
        await client.post("/auth/login", json={"email": "owner@acme.test", "password": "wrong-pass"})
    ).status_code == 401
    login = await client.post("/auth/login", json={"email": "OWNER@acme.test", "password": "password123"})
    assert login.status_code == 200

    joined = await client.post(
        "/auth/join",
        json={"email": "member@acme.test", "password": "password123", "inviteCode": s["workspace"]["inviteCode"]},
    )
    assert joined.status_code == 201 and joined.json()["user"]["role"] == "member"
    members = (await client.get("/workspaces/members", headers=auth(s))).json()
    assert {m["email"] for m in members} == {"owner@acme.test", "member@acme.test"}

    assert (await client.get("/admin/usage", headers=auth(joined.json()))).status_code == 403
    assert (await client.get("/documents")).status_code == 401
    assert (await client.get("/documents", headers={"Authorization": "Bearer nope"})).status_code == 401


async def test_upload_ingests_through_all_stages_and_answers_with_citations(client, container):
    s = await signup(client, "rag@acme.test", "RAG Team")
    doc = await upload(client, s, "msa.md", MSA, "text/markdown")
    assert doc["status"] == "ready" and doc["progress"] == 100 and doc["chunkCount"] >= 1

    stored = await container.db.documents.find_one({"title": "msa.md"})
    assert "extractedPages" not in stored, "checkpointed page text must be cleared after chunking"
    chunk = await container.db.chunks.find_one({"documentId": stored["_id"]})
    assert chunk["live"] and len(chunk["embedding"]) == 256
    assert chunk["embedText"].startswith("[msa.md - "), "contextual header must prefix the embedded text"

    events = await ask(client, s, "Does the limitation of liability apply to gross negligence?")
    sources = events["sources"][0]
    assert sources["refused"] is False and sources["cached"] is False and sources["sources"]
    answer = "".join(t["t"] for t in events["token"])
    assert "[1]" in answer
    done = events["done"][0]
    assert done["invalidCitations"] == 0 and done["queryId"]

    recorded = await container.db.queries.find_one({"question": {"$regex": "gross negligence"}})
    assert recorded["retrieved"] and recorded["retrieved"][0]["rrfScore"] > 0
    assert recorded["citations"][0]["marker"] == 1

    again = await ask(client, s, "Does the  limitation of liability apply to gross negligence?")
    assert again["sources"][0]["cached"] is True and again["done"][0]["cached"] is True

    fb = await client.post(
        f"/query/{done['queryId']}/feedback", json={"feedback": "down", "note": "x"}, headers=auth(s)
    )
    assert fb.status_code == 204
    usage = (await client.get("/admin/usage", headers=auth(s))).json()
    assert usage["queries"] == 1 and usage["capReached"] is False


async def test_new_upload_invalidates_cached_answers(client):
    s = await signup(client, "cache@acme.test", "Cache Team")
    await upload(client, s, "msa.md", MSA, "text/markdown")
    first = await ask(client, s, "termination refund")
    assert first["sources"][0]["cached"] is False
    assert (await ask(client, s, "termination refund"))["sources"][0]["cached"] is True

    await upload(client, s, "dpa.txt", b"The Processor shall notify the Controller within 72 hours.", "text/plain")
    assert (await ask(client, s, "termination refund"))["sources"][0]["cached"] is False


async def test_workspaces_are_isolated(client):
    a = await signup(client, "a@tenant.test", "Tenant A")
    b = await signup(client, "b@tenant.test", "Tenant B")
    secret = await upload(client, a, "secret.md", b"# Secret\n\nThe launch codeword is BLUEHERON.", "text/markdown")

    assert (await client.get(f"/documents/{secret['id']}", headers=auth(b))).status_code == 404
    assert (await client.get(f"/documents/{secret['id']}/file", headers=auth(b))).status_code == 404
    assert (await client.get("/documents", headers=auth(b))).json() == []

    events = await ask(client, b, "BLUEHERON launch codeword")
    assert all(src["documentId"] != secret["id"] for src in events["sources"][0]["sources"])
    assert events["sources"][0]["refused"] is True
    assert "".join(t["t"] for t in events["token"]) == REFUSAL_MESSAGE

    own = await ask(client, a, "BLUEHERON launch codeword")
    assert own["sources"][0]["sources"][0]["documentId"] == secret["id"]


async def test_delete_delists_chunks_immediately(client, container):
    s = await signup(client, "del@acme.test", "Delete Team")
    doc = await upload(client, s, "secret.md", b"# Notes\n\nThe vault combination is 4-8-15-16.", "text/markdown")
    assert (await ask(client, s, "vault combination"))["sources"][0]["sources"]

    assert (await client.delete(f"/documents/{doc['id']}", headers=auth(s))).status_code == 204
    assert (await client.get(f"/documents/{doc['id']}", headers=auth(s))).status_code == 404
    assert await container.db.chunks.count_documents({"live": True, "text": {"$regex": "vault"}}) == 0
    after = await ask(client, s, "vault combination")
    assert after["sources"][0]["refused"] is True and after["sources"][0]["cached"] is False


async def test_pdf_and_docx_uploads(client, tiny_pdf):
    from io import BytesIO

    from docx import Document

    s = await signup(client, "files@acme.test", "Files Team")
    pdf = await upload(client, s, "contract.pdf", tiny_pdf, "application/pdf")
    assert pdf["status"] == "ready" and pdf["pageCount"] == 2

    d = Document()
    d.add_heading("Remote Work Policy", level=1)
    d.add_paragraph("Employees may work remotely up to 3 days per week.")
    buf = BytesIO()
    d.save(buf)
    docx = await upload(client, s, "handbook.docx", buf.getvalue(), DOCX)
    assert docx["status"] == "ready"

    events = await ask(client, s, "How many days can employees work remotely?")
    titles = {src["documentTitle"] for src in events["sources"][0]["sources"]}
    assert "handbook.docx" in titles

    file_res = await client.get(f"/documents/{pdf['id']}/file", params={"token": s["token"]})
    assert file_res.status_code == 200 and file_res.content == tiny_pdf

    bad = await client.post("/documents", files={"file": ("x.png", b"\x89PNG", "image/png")}, headers=auth(s))
    assert bad.status_code == 415


async def test_corrupt_file_is_marked_failed_after_extract_retries(client):
    s = await signup(client, "broken@acme.test", "Broken Team")
    doc = await upload(client, s, "broken.pdf", b"%PDF-1.4 this is not really a pdf", "application/pdf")
    assert doc["status"] == "failed"
    assert doc["error"]["stage"] == "extract"


async def test_monthly_cap_blocks_further_queries(client, container):
    s = await signup(client, "cap@acme.test", "Cap Team")
    await upload(client, s, "msa.md", MSA, "text/markdown")
    await container.db.workspaces.update_one({"slug": "cap-team"}, {"$set": {"monthlyCapUsd": 0.0}})
    await container.metering.meter(s["workspace"]["id"], 0.0, 1.0)
    events = await ask(client, s, "termination")
    assert "error" in events and "cap" in events["error"][0]["message"]


async def test_query_rate_limit(client, container):
    s = await signup(client, "rate@acme.test", "Rate Team")
    original = container.settings.query_rate_limit_per_minute
    container.settings.query_rate_limit_per_minute = 2
    try:
        codes = [(await client.post("/query", json={"question": "hi"}, headers=auth(s))).status_code for _ in range(3)]
    finally:
        container.settings.query_rate_limit_per_minute = original
    assert codes == [200, 200, 429]
